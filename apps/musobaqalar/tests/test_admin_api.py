"""1-bosqich: musobaqalar modellari va ERP ichidagi boshqaruv API."""
from datetime import timedelta
from io import BytesIO

import pytest
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.musobaqalar.models import Competition, Participant, Result
from apps.notifications.models import ActivityLog

URL = '/api/v1/musobaqalar/'


def client_for(user):
    c = APIClient()
    c.force_authenticate(user=user)
    return c


@pytest.fixture
def admin(db):
    return User.objects.create_user(phone='+998908880001', password='pass1234',
                                    full_name='Admin', role=User.Role.ADMIN)


@pytest.fixture
def api(admin):
    return client_for(admin)


def make_competition(**extra):
    data = dict(name='Oktyabr matematika', grade_from=1, grade_to=11,
                registration_deadline=timezone.now() + timedelta(days=7))
    data.update(extra)
    return Competition.objects.create(**data)


def make_participant(competition, phone='+998901112233', grade=5, **extra):
    data = dict(first_name='Ali', last_name='Valiyev', address="Qorako'l")
    data.update(extra)
    return Participant.objects.create(competition=competition, phone=phone, grade=grade, **data)


# ── Modellar ─────────────────────────────────────────────────────────────────
@pytest.mark.django_db
class TestModels:

    def test_same_phone_only_once_per_competition(self):
        c = make_competition()
        make_participant(c)
        with pytest.raises(IntegrityError), transaction.atomic():
            make_participant(c)
        # Boshqa musobaqaga shu telefon bilan yozilish mumkin
        other = make_competition(name='Noyabr')
        make_participant(other)

    def test_only_one_published_competition(self):
        make_competition(status=Competition.Status.PUBLISHED)
        with pytest.raises(IntegrityError), transaction.atomic():
            make_competition(name='Ikkinchi', status=Competition.Status.PUBLISHED)

    def test_grade_range_constraint(self):
        with pytest.raises(IntegrityError), transaction.atomic():
            make_competition(grade_from=9, grade_to=5)

    def test_result_one_per_participant(self):
        p = make_participant(make_competition())
        Result.objects.create(participant=p, score=90)
        with pytest.raises(IntegrityError), transaction.atomic():
            Result.objects.create(participant=p, score=80)

    def test_registration_open_only_when_published_and_before_deadline(self):
        c = make_competition()
        assert not c.is_registration_open
        c.status = Competition.Status.PUBLISHED
        assert c.is_registration_open
        c.registration_deadline = timezone.now() - timedelta(minutes=1)
        assert not c.is_registration_open

    def test_participants_protect_competition(self):
        from django.db.models import ProtectedError
        c = make_competition()
        make_participant(c)
        with pytest.raises(ProtectedError):
            c.delete()


# ── Musobaqalar boshqaruvi ────────────────────────────────────────────────────
@pytest.mark.django_db
class TestCompetitionApi:

    def test_create_sets_creator_and_draft(self, api, admin):
        resp = api.post(f'{URL}competitions/', {
            'name': 'Oktyabr', 'grade_from': 3, 'grade_to': 7,
            'registration_deadline': (timezone.now() + timedelta(days=5)).isoformat(),
        }, format='json')
        assert resp.status_code == 201, resp.data
        assert resp.data['status'] == 'draft'
        assert resp.data['created_by_name'] == 'Admin'
        assert resp.data['next_status'] == 'published'
        assert ActivityLog.objects.filter(model_name='Competition', action='create').exists()

    def test_status_cannot_be_set_via_patch(self, api):
        c = make_competition()
        api.patch(f'{URL}competitions/{c.id}/', {'status': 'finished'}, format='json')
        c.refresh_from_db()
        assert c.status == 'draft'

    def test_invalid_grade_range_rejected(self, api):
        resp = api.post(f'{URL}competitions/', {
            'name': 'X', 'grade_from': 8, 'grade_to': 2,
            'registration_deadline': (timezone.now() + timedelta(days=5)).isoformat(),
        }, format='json')
        assert resp.status_code == 400

    def test_status_moves_forward_step_by_step(self, api):
        c = make_competition()
        for expected in ['published', 'registration_closed', 'finished']:
            resp = api.post(f'{URL}competitions/{c.id}/set-status/', {'status': expected}, format='json')
            assert resp.status_code == 200, resp.data
            assert resp.data['status'] == expected
        final = api.post(f'{URL}competitions/{c.id}/set-status/', {'status': 'draft'}, format='json')
        assert final.status_code == 400

    def test_cannot_skip_or_go_back(self, api):
        c = make_competition()
        assert api.post(f'{URL}competitions/{c.id}/set-status/', {'status': 'finished'},
                        format='json').status_code == 400
        c.status = Competition.Status.REGISTRATION_CLOSED
        c.save()
        assert api.post(f'{URL}competitions/{c.id}/set-status/', {'status': 'published'},
                        format='json').status_code == 400

    def test_publish_refused_when_another_published(self, api):
        make_competition(name='Birinchi', status=Competition.Status.PUBLISHED)
        c = make_competition(name='Ikkinchi')
        resp = api.post(f'{URL}competitions/{c.id}/set-status/', {'status': 'published'}, format='json')
        assert resp.status_code == 400
        assert 'Birinchi' in resp.data['detail']

    def test_publish_refused_after_deadline(self, api):
        c = make_competition(registration_deadline=timezone.now() - timedelta(hours=1))
        resp = api.post(f'{URL}competitions/{c.id}/set-status/', {'status': 'published'}, format='json')
        assert resp.status_code == 400

    def test_status_change_logged(self, api):
        c = make_competition()
        api.post(f'{URL}competitions/{c.id}/set-status/', {'status': 'published'}, format='json')
        log = ActivityLog.objects.filter(model_name='Competition', action='update').latest('created_at')
        assert log.changes == {'status': {'old': 'draft', 'new': 'published'}}

    def test_delete_only_empty_draft(self, api):
        empty = make_competition(name='Bo\'sh')
        assert api.delete(f'{URL}competitions/{empty.id}/').status_code == 204
        used = make_competition(name='Ishtirokchili')
        make_participant(used)
        assert api.delete(f'{URL}competitions/{used.id}/').status_code == 400

    def test_cannot_narrow_grades_below_registered(self, api):
        c = make_competition()
        make_participant(c, grade=10)
        resp = api.patch(f'{URL}competitions/{c.id}/', {'grade_to': 9}, format='json')
        assert resp.status_code == 400

    def test_counts(self, api):
        c = make_competition()
        make_participant(c, phone='+998901110001')
        make_participant(c, phone='+998901110002', status=Participant.Status.CONFIRMED)
        resp = api.get(f'{URL}competitions/{c.id}/')
        assert (resp.data['participant_count'], resp.data['confirmed_count']) == (2, 1)


# ── Ishtirokchilar jadvali ────────────────────────────────────────────────────
@pytest.mark.django_db
class TestParticipantsApi:

    @pytest.fixture
    def data(self):
        c = make_competition()
        other = make_competition(name='Boshqa')
        make_participant(c, phone='+998901110001', grade=3)
        make_participant(c, phone='+998901110002', grade=5, first_name='Sardor')
        make_participant(c, phone='+998901110003', grade=5, status=Participant.Status.CONFIRMED)
        make_participant(other, phone='+998901110004', grade=5)
        return c

    def rows(self, resp):
        return resp.data.get('results', resp.data)

    def test_filters_and_search(self, api, data):
        base = f'{URL}participants/?competition={data.id}'
        assert len(self.rows(api.get(base))) == 3
        assert len(self.rows(api.get(base + '&grade=5'))) == 2
        assert len(self.rows(api.get(base + '&status=new'))) == 2
        assert len(self.rows(api.get(base + '&status=confirmed'))) == 1
        assert [r['phone'] for r in self.rows(api.get(base + '&search=Sardor&grade=5'))] == ['+998901110002']
        assert [r['phone'] for r in self.rows(api.get(base + '&search=110001'))] == ['+998901110001']

    def test_confirm(self, api, admin, data):
        p = Participant.objects.get(phone='+998901110001')
        resp = api.post(f'{URL}participants/{p.id}/confirm/')
        assert resp.status_code == 200
        p.refresh_from_db()
        assert (p.status, p.confirmed_by) == ('confirmed', admin)
        assert p.confirmed_at is not None
        assert api.post(f'{URL}participants/{p.id}/confirm/').status_code == 400

    def test_participants_read_only(self, api, data):
        p = Participant.objects.first()
        assert api.patch(f'{URL}participants/{p.id}/', {'grade': 9}, format='json').status_code == 405
        assert api.delete(f'{URL}participants/{p.id}/').status_code == 405

    def test_excel_export(self, api, data):
        resp = api.get(f'{URL}competitions/{data.id}/export/')
        assert resp.status_code == 200
        ws = load_workbook(BytesIO(resp.content)).active
        assert ws.cell(row=2, column=2).value == 'Familya'
        phones = [ws.cell(row=r, column=5).value for r in range(3, ws.max_row + 1)]
        assert sorted(phones) == ['+998901110001', '+998901110002', '+998901110003']


# ── Ruxsatlar ─────────────────────────────────────────────────────────────────
@pytest.mark.django_db
@pytest.mark.parametrize('role', [User.Role.TEACHER, User.Role.FINANCE, User.Role.STUDENT, User.Role.PARENT])
def test_only_admin_manages_competitions(role):
    user = User.objects.create_user(phone='+998908880009', password='pass1234', full_name='X', role=role)
    c = client_for(user)
    assert c.get(f'{URL}competitions/').status_code == 403
    assert c.get(f'{URL}participants/').status_code == 403


@pytest.mark.django_db
def test_anonymous_denied():
    assert APIClient().get(f'{URL}competitions/').status_code == 401
