"""
Musobaqalar — ERP ichidagi boshqaruv API (faqat administrator).
Public sayt uchun endpointlar alohida: public_views.py (2-bosqich).
"""
from django.db import transaction
from django.db.models import Count, F, Q
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.response import Response

from apps.accounts.permissions import IsAdmin
from apps.notifications.models import ActivityLog
from apps.notifications.views import diff_fields, log_activity
from apps.students.export_views import _add_title, _response, _style_header, _style_row

from .models import Competition, Participant, Result
from .serializers import (
    CompetitionSerializer,
    ParticipantSerializer,
    ResultInputSerializer,
    SetStatusSerializer,
)

# Ball faqat ro'yxat yopilgandan keyin kiritiladi; public saytda faqat "yakunlangan"da ko'rinadi
RESULT_STATUSES = (Competition.Status.REGISTRATION_CLOSED, Competition.Status.FINISHED)


class CompetitionViewSet(viewsets.ModelViewSet):
    """
    /api/v1/musobaqalar/competitions/          — ro'yxat, yaratish (?status=)
    /api/v1/musobaqalar/competitions/<id>/     — ko'rish, tahrirlash, o'chirish (faqat qoralama)
    POST .../<id>/set-status/ {"status": ...}   — holatni keyingisiga o'tkazish
    GET  .../<id>/export/                       — ishtirokchilar Excel
    """
    serializer_class = CompetitionSerializer
    permission_classes = [IsAdmin]
    filterset_fields = ['status']
    search_fields = ['name']
    ordering = ['-created_at']

    def get_queryset(self):
        return (Competition.objects.select_related('created_by')
                .annotate(participant_count=Count('participants'),
                          confirmed_count=Count('participants',
                                                filter=Q(participants__status=Participant.Status.CONFIRMED))))

    def perform_create(self, serializer):
        competition = serializer.save(created_by=self.request.user)
        log_activity(self.request.user, ActivityLog.Action.CREATE, 'Competition',
                     competition.pk, str(competition), request=self.request)

    def perform_update(self, serializer):
        fields = ('name', 'grade_from', 'grade_to', 'location', 'registration_deadline',
                  'competition_date', 'show_full_names')
        before = {f: getattr(serializer.instance, f) for f in fields}
        competition = serializer.save()
        log_activity(self.request.user, ActivityLog.Action.UPDATE, 'Competition',
                     competition.pk, str(competition), changes=diff_fields(before, competition, fields),
                     request=self.request)

    def destroy(self, request, *args, **kwargs):
        competition = self.get_object()
        if competition.status != Competition.Status.DRAFT or competition.participants.exists():
            return Response({'detail': "Faqat ishtirokchisi yo'q qoralama musobaqani o'chirish mumkin."},
                            status=status.HTTP_400_BAD_REQUEST)
        log_activity(request.user, ActivityLog.Action.DELETE, 'Competition',
                     competition.pk, str(competition), request=request)
        competition.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)

    @action(detail=True, methods=['post'], url_path='set-status')
    def set_status(self, request, pk=None):
        competition = self.get_object()
        serializer = SetStatusSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        new_status = serializer.validated_data['status']
        expected = Competition.NEXT_STATUS.get(competition.status)
        if new_status != expected:
            label = Competition.Status(expected).label if expected else None
            return Response({'detail': (f"Holat tartib bilan o'zgaradi: keyingisi — \"{label}\"."
                                        if label else "Yakunlangan musobaqa holati o'zgarmaydi.")},
                            status=status.HTTP_400_BAD_REQUEST)
        if new_status == Competition.Status.PUBLISHED:
            if competition.registration_deadline <= timezone.now():
                return Response({'detail': "Ro'yxatdan o'tish muddati o'tib ketgan — avval muddatni yangilang."},
                                status=status.HTTP_400_BAD_REQUEST)
            other = Competition.objects.filter(status=Competition.Status.PUBLISHED).exclude(pk=competition.pk).first()
            if other:
                return Response({'detail': (f"\"{other.name}\" allaqachon e'lon qilingan — bir vaqtda bitta "
                                            "musobaqa e'lon qilinadi. Avval uning ro'yxatini yoping.")},
                                status=status.HTTP_400_BAD_REQUEST)
        old = competition.status
        competition.status = new_status
        competition.save(update_fields=['status', 'updated_at'])
        log_activity(request.user, ActivityLog.Action.UPDATE, 'Competition', competition.pk, str(competition),
                     changes={'status': {'old': old, 'new': new_status}}, request=request)
        return Response(self.get_serializer(self.get_queryset().get(pk=competition.pk)).data)

    @action(detail=True, methods=['post'])
    def rank(self, request, pk=None):
        """POST .../<id>/rank/ — o'rinlarni ball bo'yicha avtomatik hisoblash (har bir sinf ichida,
        sinfsizlar alohida umumiy guruhda). Teng ball — bir xil o'rin: 1, 1, 3 ..."""
        competition = self.get_object()
        if competition.status not in RESULT_STATUSES:
            return Response({'detail': "Natijalar ro'yxat yopilgandan keyin kiritiladi."},
                            status=status.HTTP_400_BAD_REQUEST)
        results = list(Result.objects.filter(participant__competition=competition)
                       .select_related('participant').order_by('-score'))
        groups = {}
        for r in results:
            groups.setdefault(r.participant.grade, []).append(r)
        updated = 0
        with transaction.atomic():
            for rows in groups.values():
                place, prev_score = 0, None
                for i, r in enumerate(rows, 1):
                    if r.score != prev_score:
                        place, prev_score = i, r.score
                    if r.place != place:
                        r.place = place
                        r.save(update_fields=['place', 'entered_at'])
                        updated += 1
        log_activity(request.user, ActivityLog.Action.UPDATE, 'Competition', competition.pk, str(competition),
                     changes={'rank': {'old': None, 'new': f'{len(results)} ta natija'}}, request=request)
        return Response({'ranked': len(results), 'updated': updated})

    @action(detail=True, methods=['get'])
    def export(self, request, pk=None):
        competition = self.get_object()
        rows = (competition.participants.select_related('result', 'confirmed_by')
                .order_by(F('grade').asc(nulls_last=True), 'last_name', 'first_name'))
        wb = Workbook()
        ws = wb.active
        ws.title = 'Ishtirokchilar'
        ws.freeze_panes = 'A3'
        headers = ['#', 'Familya', 'Ism', 'Sinf', 'Telefon', 'Yashash manzili', 'Holati',
                   'Tasdiqlagan', "Ro'yxatdan o'tgan", 'Ball', "O'rin"]
        _add_title(ws, f"{competition.name} — ishtirokchilar", len(headers))
        _style_header(ws, headers)
        for i, p in enumerate(rows, 1):
            result = getattr(p, 'result', None)
            _style_row(ws, i + 2, [
                i, p.last_name, p.first_name, p.grade or '', p.phone, p.address, p.get_status_display(),
                p.confirmed_by.full_name if p.confirmed_by else '—',
                timezone.localtime(p.registered_at).strftime('%d.%m.%Y %H:%M'),
                result.score if result else '', (result.place or '') if result else '',
            ], i % 2 == 0)
        for col, w in enumerate([5, 18, 16, 7, 16, 32, 13, 18, 17, 8, 8], 1):
            ws.column_dimensions[get_column_letter(col)].width = w
        return _response(wb, f"musobaqa_ishtirokchilar_{timezone.localdate():%Y%m%d}.xlsx")


class ParticipantViewSet(viewsets.ReadOnlyModelViewSet):
    """
    GET  /api/v1/musobaqalar/participants/?competition=<id>&grade=&status=&search=
    POST /api/v1/musobaqalar/participants/<id>/confirm/  — operator qo'ng'iroqdan keyin tasdiqlaydi
    """
    serializer_class = ParticipantSerializer
    permission_classes = [IsAdmin]
    filterset_fields = ['competition', 'grade', 'status']
    search_fields = ['first_name', 'last_name', 'phone', 'address']
    ordering_fields = ['registered_at', 'grade', 'last_name']
    ordering = ['-registered_at']

    def get_queryset(self):
        return Participant.objects.select_related('competition', 'confirmed_by', 'result')

    @action(detail=True, methods=['post'])
    def confirm(self, request, pk=None):
        participant = self.get_object()
        if participant.status == Participant.Status.CONFIRMED:
            return Response({'detail': 'Bu ishtirokchi allaqachon tasdiqlangan.'},
                            status=status.HTTP_400_BAD_REQUEST)
        participant.status = Participant.Status.CONFIRMED
        participant.confirmed_by = request.user
        participant.confirmed_at = timezone.now()
        participant.save(update_fields=['status', 'confirmed_by', 'confirmed_at'])
        log_activity(request.user, ActivityLog.Action.UPDATE, 'Participant', participant.pk, str(participant),
                     changes={'status': {'old': Participant.Status.NEW, 'new': Participant.Status.CONFIRMED}},
                     request=request)
        return Response(self.get_serializer(participant).data)

    @action(detail=True, methods=['post'])
    def result(self, request, pk=None):
        """POST .../participants/<id>/result/ {"score": 87, "place": 2}
        Ball kiritish/o'zgartirish (o'rin ixtiyoriy — "rank" bilan avtomatik ham hisoblanadi).
        {"score": null} — natijani o'chirish."""
        participant = self.get_object()
        if participant.competition.status not in RESULT_STATUSES:
            return Response({'detail': "Natijalar ro'yxat yopilgandan keyin kiritiladi."},
                            status=status.HTTP_400_BAD_REQUEST)
        serializer = ResultInputSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        score = serializer.validated_data.get('score')
        existing = Result.objects.filter(participant=participant).first()
        old = {'score': existing.score, 'place': existing.place} if existing else {'score': None, 'place': None}
        if score is None:
            if existing:
                existing.delete()
            new = {'score': None, 'place': None}
        else:
            place = serializer.validated_data.get('place')
            Result.objects.update_or_create(
                participant=participant,
                defaults={'score': score, 'place': place, 'entered_by': request.user},
            )
            new = {'score': score, 'place': place}
        changes = {k: {'old': str(old[k]) if old[k] is not None else None,
                       'new': str(new[k]) if new[k] is not None else None}
                   for k in ('score', 'place') if old[k] != new[k]}
        log_activity(request.user, ActivityLog.Action.UPDATE, 'Participant', participant.pk, str(participant),
                     changes=changes or None, request=request)
        participant = self.get_queryset().get(pk=participant.pk)
        return Response(self.get_serializer(participant).data)

